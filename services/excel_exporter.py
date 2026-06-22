"""
Excel Export Service
Generates admin review Excel report after appointment booking
Required by assignment: "Export to Excel for admin review"
"""

import os
from datetime import datetime
from typing import cast

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet


def export_appointment_to_excel(appointment_data: dict, output_dir: str = "files") -> str:
    """
    Generate a formatted Excel report for admin review.
    Called after appointment is confirmed (step 7).

    Args:
        appointment_data: Dict with all appointment details
        output_dir: Directory to save the file

    Returns:
        Path to the generated Excel file
    """
    os.makedirs(output_dir, exist_ok=True)

    wb = openpyxl.Workbook()
    ws: Worksheet = cast(Worksheet, wb.active)
    ws.title = "Appointment Report"

    # ── Styles ──────────────────────────────────────────────
    header_font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1a5cdf")
    label_font = Font(name="Calibri", bold=True, size=11, color="1a2840")
    value_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    alt_fill = PatternFill("solid", fgColor="F0F4F8")

    # ── Title ────────────────────────────────────────────────
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = "🏥 MediBook — Appointment Report"
    title_cell.font = Font(name="Calibri", bold=True, size=16, color="1a2840")
    title_cell.alignment = center_align
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:D2")
    gen_cell = ws["A2"]
    gen_cell.value = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    gen_cell.font = Font(name="Calibri", size=10, color="64748B", italic=True)
    gen_cell.alignment = center_align
    ws.row_dimensions[2].height = 20

    ws.append([])  # blank row

    # ── Section header ────────────────────────────────────────
    def section_header(row_idx, title):
        ws.merge_cells(f"A{row_idx}:D{row_idx}")
        cell = ws[f"A{row_idx}"]
        cell.value = title
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        ws.row_dimensions[row_idx].height = 24

    def data_row(label, value, row_idx, alt=False):
        fill = alt_fill if alt else PatternFill("solid", fgColor="FFFFFF")
        for col in ["A", "B", "C", "D"]:
            ws[f"{col}{row_idx}"].fill = fill
            ws[f"{col}{row_idx}"].border = thin_border

        ws[f"A{row_idx}"].value = label
        ws[f"A{row_idx}"].font = label_font
        ws[f"A{row_idx}"].alignment = left_align

        ws.merge_cells(f"B{row_idx}:D{row_idx}")
        ws[f"B{row_idx}"].value = str(value) if value else "—"
        ws[f"B{row_idx}"].font = value_font
        ws[f"B{row_idx}"].alignment = left_align
        ws.row_dimensions[row_idx].height = 22

    # ── Section 1: Patient Info ──────────────────────────────
    r = 4
    section_header(r, "👤 Patient Information")
    rows_p = [
        ("Patient Name", appointment_data.get("patient_name")),
        ("Patient ID", appointment_data.get("patient_id")),
        ("Patient Type", appointment_data.get("patient_type", "New").title()),
        ("Date of Birth", appointment_data.get("dob")),
        ("Email", appointment_data.get("email")),
        ("Phone", appointment_data.get("phone")),
    ]
    for i, (label, val) in enumerate(rows_p):
        data_row(label, val, r + 1 + i, alt=(i % 2 == 1))
    r += len(rows_p) + 2

    # ── Section 2: Appointment ───────────────────────────────
    section_header(r, "📅 Appointment Details")
    rows_a = [
        ("Appointment ID", appointment_data.get("appointment_id")),
        ("Doctor", appointment_data.get("preferred_doctor")),
        ("Location", appointment_data.get("location")),
        ("Date", appointment_data.get("appointment_date")),
        ("Time", appointment_data.get("selected_time")),
        ("Duration", f"{appointment_data.get('appointment_duration', 30)} minutes"),
        (
            "Booking Confirmed",
            "✅ Yes" if appointment_data.get("booking_confirmed") else "⏳ Pending",
        ),
    ]
    for i, (label, val) in enumerate(rows_a):
        data_row(label, val, r + 1 + i, alt=(i % 2 == 1))
    r += len(rows_a) + 2

    # ── Section 3: Insurance ─────────────────────────────────
    section_header(r, "🏥 Insurance Information")
    rows_i = [
        ("Insurance Carrier", appointment_data.get("insurance_carrier")),
        ("Member ID", appointment_data.get("insurance_member_id")),
        ("Group ID", appointment_data.get("insurance_group_id")),
    ]
    for i, (label, val) in enumerate(rows_i):
        data_row(label, val, r + 1 + i, alt=(i % 2 == 1))
    r += len(rows_i) + 2

    # ── Section 4: Reminders ─────────────────────────────────
    section_header(r, "🔔 Reminders Scheduled")
    appt_dt = f"{appointment_data.get('appointment_date', '')} {appointment_data.get('selected_time', '')}"
    reminder_rows = [
        ("Reminder 1 (48h before)", "General reminder — appointment coming up"),
        ("Reminder 2 (24h before)", "Action: Have you filled your intake forms?"),
        ("Reminder 3 (1h before)", "Action: Is your visit confirmed? Reply YES/NO"),
    ]
    for i, (label, val) in enumerate(reminder_rows):
        data_row(label, val, r + 1 + i, alt=(i % 2 == 1))

    # ── Column widths ────────────────────────────────────────
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20

    # ── Save ─────────────────────────────────────────────────
    appt_id = appointment_data.get("appointment_id", "unknown")
    filename = f"appointment_{appt_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(output_dir, filename)
    wb.save(filepath)

    return filepath
